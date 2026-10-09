#!/usr/bin/env python3
# mist_ap_ssid_channel_report.py
#
# For every AP in the org, shows which SSIDs it publishes on which channel,
# and which Device Profile is applied to it.
#
# How the SSIDs are worked out:
#   - /sites/{site_id}/wlans/derived gives every WLAN that reaches the site
#     (org WLAN Templates + site WLANs)
#   - a WLAN from a Template with "filter by Device Profile" only reaches APs
#     whose Device Profile is in that Template's list
#   - a WLAN set to "apply to specific APs" only reaches those APs
#   - each WLAN goes out on the bands in its "bands" setting
# The channel, width, power and clients come from the AP's live radio stats.
# The live stats also say how many SSIDs each radio is really broadcasting
# (num_wlans), which is used to check the worked-out list.
#
# Output: an Excel report, one row per AP plus one row per AP / band / SSID.
#
# Usage:
#   Edit mist_ap_ssid_channel_report.ini (same folder as this script), then:
#   python3 mist_ap_ssid_channel_report.py
#   python3 mist_ap_ssid_channel_report.py other_org.ini     (optional: use a different .ini)

import os
import sys
import time
import datetime
import configparser
import re
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ---------------- read settings from the .ini file ----------------

if len(sys.argv) > 1:
    INI_FILE = sys.argv[1]
else:
    INI_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "mist_ap_ssid_channel_report.ini")

if not os.path.exists(INI_FILE):
    print("Can't find settings file: " + INI_FILE)
    print("Copy mist_ap_ssid_channel_report.ini.example to mist_ap_ssid_channel_report.ini and fill it in.")
    sys.exit(1)

config = configparser.ConfigParser()
config.read(INI_FILE)

API_HOST = config.get("mist", "api_host", fallback="https://api.eu.mist.com").rstrip("/")
API_TOKEN = config.get("mist", "api_token", fallback="").strip()
ORG_ID = config.get("mist", "org_id", fallback="").strip()

SITE_NAME_FILTER = config.get("run", "site_name_filter", fallback="").strip()
PAUSE_BETWEEN_CALLS = config.getfloat("run", "pause_between_calls", fallback=0)
REPORT_FOLDER = os.path.expanduser(config.get("run", "report_folder", fallback=".").strip() or ".")

# Optional [network] section, for company HTTPS proxies such as Zscaler
USE_SYSTEM_CERTS = config.getboolean("network", "use_system_certs", fallback=True)
CA_BUNDLE = os.path.expanduser(config.get("network", "ca_bundle", fallback="").strip())
VERIFY_SSL = config.getboolean("network", "verify_ssl", fallback=True)
HTTP_PROXY = config.get("network", "http_proxy", fallback="").strip()
HTTPS_PROXY = config.get("network", "https_proxy", fallback="").strip()

# -----------------------------------------------------------------

if API_TOKEN == "" or API_TOKEN.startswith("PASTE"):
    print("Put your Mist API token in " + INI_FILE + " ([mist] api_token)")
    sys.exit(1)
if ORG_ID == "" or ORG_ID.startswith("PASTE"):
    print("Put the org ID in " + INI_FILE + " ([mist] org_id)")
    sys.exit(1)

# ---------- HTTPS certificates and proxies ----------
# On company networks a proxy (Zscaler, Netskope, Palo Alto etc.) often re-signs
# HTTPS traffic with the company's own root certificate. That certificate is in the
# Windows/macOS certificate store, but not in the bundle Python uses by default,
# which gives "CERTIFICATE_VERIFY_FAILED: unable to get local issuer certificate".
#   use_system_certs = true -> trust the operating system's certificate store (needs truststore)
#   ca_bundle = <path>      -> or point at the company root certificate (.pem) instead
#   verify_ssl = false      -> last resort only: turns certificate checking off

if USE_SYSTEM_CERTS and CA_BUNDLE == "" and VERIFY_SSL:
    try:
        import truststore
        truststore.inject_into_ssl()
    except ImportError:
        print("Note: 'truststore' isn't installed, so the OS certificate store isn't used.")
        print("      If you get CERTIFICATE_VERIFY_FAILED, run: pip install -r requirements.txt")
        print("")

session = requests.Session()
session.headers.update({"Authorization": "Token " + API_TOKEN})
if CA_BUNDLE != "":
    if not os.path.exists(CA_BUNDLE):
        print("ca_bundle file not found: " + CA_BUNDLE)
        sys.exit(1)
    session.verify = CA_BUNDLE
    # requests lets these environment variables override session.verify, so set them too
    os.environ["REQUESTS_CA_BUNDLE"] = CA_BUNDLE
    os.environ["CURL_CA_BUNDLE"] = CA_BUNDLE
elif not VERIFY_SSL:
    print("WARNING: certificate checking is OFF ([network] verify_ssl = false) - only use this as a last resort.")
    print("")
    session.verify = False
    requests.packages.urllib3.disable_warnings(requests.packages.urllib3.exceptions.InsecureRequestWarning)
if HTTP_PROXY != "":
    session.proxies["http"] = HTTP_PROXY
if HTTPS_PROXY != "":
    session.proxies["https"] = HTTPS_PROXY

api_calls = 0


def api_get(path, params=None):
    """GET from the Mist API. Waits and retries on rate limiting (429) and
    short outages (5xx). Stops the script on anything else, so the report is
    never written with sites missing."""
    global api_calls
    url = API_HOST + "/api/v1" + path
    tries = 0
    while True:
        tries = tries + 1
        try:
            api_calls = api_calls + 1
            resp = session.get(url, params=params, timeout=60)
        except requests.exceptions.SSLError as e:
            print("")
            print("HTTPS certificate check failed connecting to " + API_HOST)
            print("")
            print("This usually means a company proxy (e.g. Zscaler) is inspecting HTTPS traffic. To fix it, either:")
            print("  1. pip install -r requirements.txt   (installs truststore, so Windows/macOS")
            print("     certificates are trusted - leave use_system_certs = true in the .ini), or")
            print("  2. get your company's root certificate as a .pem file and set")
            print("     ca_bundle = C:\\path\\to\\company-root.pem   under [network] in the .ini")
            print("")
            print("Details: " + str(e))
            sys.exit(1)
        except requests.exceptions.ProxyError as e:
            print("")
            print("Couldn't reach the proxy - check http_proxy / https_proxy under [network] in the .ini.")
            print("Details: " + str(e))
            sys.exit(1)
        except (requests.exceptions.ConnectionError, requests.exceptions.Timeout) as e:
            if tries < 4:
                print("  connection problem, retrying in 10s... (" + str(e)[:80] + ")")
                time.sleep(10)
                continue
            print("")
            print("Couldn't connect to " + API_HOST + " - check the api_host setting and network/proxy access.")
            print("Details: " + str(e))
            sys.exit(1)

        if resp.status_code == 429:
            print("  rate limited, waiting 60s...")
            time.sleep(60)
            continue
        if resp.status_code >= 500 and tries < 4:
            print("  HTTP " + str(resp.status_code) + " from Mist, retrying in 10s...")
            time.sleep(10)
            continue
        if resp.status_code != 200:
            print("")
            print("STOPPED: HTTP " + str(resp.status_code) + " from GET " + path)
            print("         " + resp.text[:300])
            if resp.status_code in (401, 403):
                print("         Check the API token has access to this org.")
            print("No report written, so you don't get one with sites missing.")
            sys.exit(1)
        if PAUSE_BETWEEN_CALLS > 0:
            time.sleep(PAUSE_BETWEEN_CALLS)
        return resp.json()


def api_get_all(path, params=None):
    """GET every page of a list (Mist uses limit/page for these lists)."""
    params = dict(params or {})
    params["limit"] = 1000
    page = 1
    items = []
    while True:
        params["page"] = page
        batch = api_get(path, params)
        if not isinstance(batch, list):
            return items
        items.extend(batch)
        if len(batch) < 1000:
            return items
        page = page + 1


BANDS = ["24", "5", "6"]
BAND_NAMES = {"24": "2.4GHz", "5": "5GHz", "6": "6GHz"}

run_now = datetime.datetime.now(datetime.timezone.utc)
run_time = run_now.strftime("%Y-%m-%d %H:%M:%S UTC")
print("Run started " + run_time)

# ---------- 1. org-level data ----------

org = api_get("/orgs/" + ORG_ID)
org_name = org.get("name", ORG_ID) or ORG_ID
org_file_part = re.sub(r"[^A-Za-z0-9_-]+", "-", org_name).strip("-") or "org"
print("Org " + org_name + " (" + ORG_ID + ") on " + API_HOST)
print("")

sites = api_get_all("/orgs/" + ORG_ID + "/sites")
if SITE_NAME_FILTER != "":
    sites = [s for s in sites if SITE_NAME_FILTER.lower() in s.get("name", "").lower()]
sites.sort(key=lambda s: s.get("name", "").lower())

templates = api_get_all("/orgs/" + ORG_ID + "/templates")
templates_by_id = {}
for t in templates:
    templates_by_id[t["id"]] = t

deviceprofiles = api_get_all("/orgs/" + ORG_ID + "/deviceprofiles", {"type": "ap"})
deviceprofiles_by_id = {}
for dp in deviceprofiles:
    deviceprofiles_by_id[dp["id"]] = dp

print("Sites to check: " + str(len(sites)) + "   WLAN Templates: " + str(len(templates))
      + "   AP Device Profiles: " + str(len(deviceprofiles)))
print("")


def wlan_bands(wlan):
    """Bands a WLAN goes out on, as a list of "24" / "5" / "6"."""
    bands = wlan.get("bands")
    if not bands:
        old = wlan.get("band", "")
        if old == "both" or old == "":
            bands = ["24", "5"]
        else:
            bands = [old]
    return [str(b) for b in bands]


def wlan_reaches_ap(wlan, ap):
    """Does this WLAN (from the site's derived list) reach this AP?
    Returns (True/False, note). The note is "wxtag" when the WLAN is
    targeted by WxLAN tags, which this script can't resolve."""
    if wlan.get("enabled") is False:
        return False, ""
    # Template that only goes to some Device Profiles
    template = templates_by_id.get(wlan.get("template_id") or "")
    if template is not None and template.get("filter_by_deviceprofile"):
        if (ap.get("deviceprofile_id") or "") not in (template.get("deviceprofile_ids") or []):
            return False, ""
    # WLAN applied to specific APs or WxLAN tags, not the whole site
    apply_to = wlan.get("apply_to") or "site"
    if apply_to == "aps":
        if ap.get("id") not in (wlan.get("ap_ids") or []):
            return False, ""
    if apply_to == "wxtags":
        return True, "wxtag"
    return True, ""


def radio_limits(ap, band):
    """Allowed channels, power range and antenna gain for a band, from the AP's
    effective radio config (its Device Profile plus any AP-level override - the
    AP stats carry this as radio_config). Returns (channels, pmin, pmax, ant_gain)."""
    rc = ap.get("radio_config") or {}
    bc = rc.get("band_" + band) or {}
    return bc.get("channels") or [], bc.get("power_min"), bc.get("power_max"), rc.get("ant_gain_" + band)


# ---------- 2. per site: APs, AP config, derived WLANs ----------

ap_rows = []
detail_rows = []
check_rows = []
dp_ap_count = {}
total_aps = 0
aps_connected = 0
aps_ssid_mismatch = 0
aps_outside_limits = 0
aps_power_capped = 0
aps_wxtag = 0
start = time.time()

for site_num, site in enumerate(sites, 1):
    site_id = site["id"]
    site_name = site.get("name", site_id)

    aps = api_get_all("/sites/" + site_id + "/stats/devices", {"type": "ap"})
    if len(aps) == 0:
        print("[" + str(site_num) + "/" + str(len(sites)) + "] " + site_name + ": no APs")
        continue

    derived = api_get("/sites/" + site_id + "/wlans/derived")
    if not isinstance(derived, list):
        derived = []

    for ap in aps:
        total_aps = total_aps + 1
        connected = (ap.get("status") == "connected")
        if connected:
            aps_connected = aps_connected + 1

        dp_id = ap.get("deviceprofile_id") or ""
        dp_name = ap.get("deviceprofile_name") or (deviceprofiles_by_id.get(dp_id, {}).get("name")) or ("(none)" if dp_id == "" else dp_id)
        dp_ap_count[dp_id] = dp_ap_count.get(dp_id, 0) + 1

        # SSIDs this AP should publish, per band
        ssids_by_band = {"24": [], "5": [], "6": []}
        has_wxtag = False
        for wlan in derived:
            reaches, note = wlan_reaches_ap(wlan, ap)
            if not reaches:
                continue
            if note == "wxtag":
                has_wxtag = True
            for band in wlan_bands(wlan):
                if band in ssids_by_band:
                    ssids_by_band[band].append((wlan, note))
        if has_wxtag:
            aps_wxtag = aps_wxtag + 1

        # Live radios, keyed by the band they are really working on
        # (a dual-band radio reports its band in "usage", e.g. band_24 running as 5GHz)
        live = {"24": [], "5": [], "6": []}
        for key, stat in (ap.get("radio_stat") or {}).items():
            if not isinstance(stat, dict):
                continue
            usage = str(stat.get("usage") or key.replace("band_", ""))
            if usage in live:
                live[usage].append(stat)

        uptime_days = round((ap.get("uptime") or 0) / 86400, 1) if connected else ""
        row = [
            site_name, ap.get("name", "") or ap.get("mac", ""), ap.get("mac", ""), ap.get("model", ""),
            ap.get("serial", ""), ap.get("version", ""), ap.get("status", ""), uptime_days, dp_name,
        ]
        notes = []
        limit_notes = []
        capped_notes = []

        for band in BANDS:
            # A radio is on if it has a channel and isn't disabled (a disabled radio still reports a channel)
            radios = [r for r in live[band] if r.get("channel") and not r.get("disabled")]
            wlans_here = ssids_by_band[band]
            ssid_names = [w.get("ssid", "") + (" (wxtag)" if n == "wxtag" else "") for w, n in wlans_here]
            ssid_names_sorted = sorted(set(ssid_names), key=str.lower)

            if len(radios) == 1:
                # Keep numbers as numbers, so Excel can sort and filter them
                channel = radios[0].get("channel")
                width = radios[0].get("bandwidth", "")
                power = radios[0].get("power", "")
            else:
                # Two radios on the same band (e.g. a dual-5GHz AP) - show both
                channel = " / ".join([str(r.get("channel")) for r in radios])
                width = " / ".join([str(r.get("bandwidth", "")) for r in radios])
                power = " / ".join([str(r.get("power", "")) for r in radios])
            clients = sum([r.get("num_clients") or 0 for r in radios]) if radios else ""
            live_count = sum([r.get("num_wlans") or 0 for r in radios]) if radios else ""

            band_disabled = bool(((ap.get("radio_config") or {}).get("band_" + band) or {}).get("disabled"))
            if not connected:
                shown = ", ".join(ssid_names_sorted)
            elif len(live[band]) == 0 and len(ssid_names) == 0:
                # The AP has no radio for this band (e.g. no 6GHz radio) and nothing is configured on it
                shown = ""
            elif not radios and band_disabled:
                # Turned off on purpose in the Device Profile / AP config - nothing goes out on this band
                shown = "(radio disabled)"
            elif not radios:
                shown = "(radio off)" if len(ssid_names) == 0 else "(radio off) " + ", ".join(ssid_names_sorted)
                if len(ssid_names) > 0:
                    notes.append(BAND_NAMES[band] + ": " + str(len(ssid_names)) + " SSID(s) configured but radio off")
            else:
                shown = ", ".join(ssid_names_sorted)
                if not has_wxtag and live_count != len(wlans_here):
                    notes.append(BAND_NAMES[band] + ": config " + str(len(wlans_here)) + ", live " + str(live_count))

            allowed, pmin, pmax, ant_gain = radio_limits(ap, band)
            row = row + [channel, width, power, ant_gain if ant_gain is not None else "", clients, shown,
                         len(wlans_here), live_count]

            # Channel / power outside what the Device Profile (or AP override) allows
            #   - channel not in the allowed list, or power above power_max: flagged
            #   - power below power_min: only noted - Mist turns the power down to stay within
            #     the legal EIRP limit (antenna gain, or a low-power channel such as 149-165 in the UK)
            if connected and radios:
                for r in radios:
                    ch = r.get("channel")
                    pw = r.get("power")
                    if allowed and ch not in allowed:
                        limit_notes.append(BAND_NAMES[band] + " ch " + str(ch) + " not in allowed list")
                    if pw is not None and pmax is not None and pw > pmax:
                        limit_notes.append(BAND_NAMES[band] + " power " + str(pw) + " above max " + str(pmax))
                    if pw is not None and pmin is not None and pw < pmin:
                        capped_notes.append(BAND_NAMES[band] + " power " + str(pw) + " below min " + str(pmin)
                                            + " (ant gain " + str(ant_gain if ant_gain is not None else "?")
                                            + " dBi, ch " + str(ch) + ")")

            # Detail rows: one per SSID on this band
            for wlan, note in wlans_here:
                template = templates_by_id.get(wlan.get("template_id") or "", {})
                source_name = ("Template: " + template.get("name", "")) if template else "Site WLAN"
                auth = (wlan.get("auth") or {}).get("type", "")
                detail_rows.append([
                    site_name, row[1], ap.get("mac", ""), ap.get("model", ""), ap.get("status", ""), dp_name,
                    BAND_NAMES[band], channel if connected else "", width if connected else "",
                    power if connected else "",
                    "yes" if (connected and radios) else ("AP not connected" if not connected
                                                          else ("disabled in config" if band_disabled else "radio off")),
                    wlan.get("ssid", ""), wlan.get("id", ""), source_name, auth, wlan.get("vlan_id", ""),
                    "yes" if wlan.get("hide_ssid") else "no",
                    "WxLAN tag - may not apply" if note == "wxtag" else "",
                ])

        if not connected:
            ssid_check = "AP not connected"
        elif has_wxtag:
            ssid_check = "Not checked (WxLAN tag WLANs)"
        elif notes:
            ssid_check = "; ".join(notes)
            aps_ssid_mismatch = aps_ssid_mismatch + 1
        else:
            ssid_check = "OK"
        if limit_notes:
            aps_outside_limits = aps_outside_limits + 1
        if capped_notes:
            aps_power_capped = aps_power_capped + 1
        row = row + [ssid_check, "; ".join(limit_notes), "; ".join(capped_notes)]
        ap_rows.append(row)

        if connected and (notes or limit_notes):
            check_rows.append([site_name, row[1], ap.get("mac", ""), ap.get("model", ""), dp_name,
                               "; ".join(notes), "; ".join(limit_notes)])

    elapsed = time.time() - start
    eta = elapsed / site_num * (len(sites) - site_num)
    print("[" + str(site_num) + "/" + str(len(sites)) + "] " + site_name + ": " + str(len(aps)) + " APs, "
          + str(len(derived)) + " WLANs   (ETA " + str(int(eta // 60)) + "m " + str(int(eta % 60)) + "s)")

# ---------- 3. Excel report ----------

if SITE_NAME_FILTER != "":
    name_part = SITE_NAME_FILTER
elif len(sites) == 1:
    name_part = sites[0].get("name", "site")
else:
    name_part = "ALL-SITES"
name_part = re.sub(r"[^A-Za-z0-9_-]+", "-", name_part).strip("-")
report_file = os.path.join(REPORT_FOLDER, "ap_ssid_channel_report_" + org_file_part + "_" + name_part + "_"
                           + run_now.strftime("%Y-%m-%d_%H%M") + "UTC.xlsx")

print("")
print("Writing Excel report (" + str(len(ap_rows)) + " APs, " + str(len(detail_rows)) + " SSID rows)...")

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="1F3864")
band_fills = {"24": PatternFill("solid", fgColor="2E75B6"), "5": PatternFill("solid", fgColor="548235"),
              "6": PatternFill("solid", fgColor="7030A0")}
amber_fill = PatternFill("solid", fgColor="FFE699")
red_fill = PatternFill("solid", fgColor="F8CBAD")
grey_fill = PatternFill("solid", fgColor="D9D9D9")
title_font = Font(bold=True, size=14)


def add_sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for r in rows:
        ws.append(r)
    for i in range(len(widths)):
        ws.column_dimensions[get_column_letter(i + 1)].width = widths[i]
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:" + get_column_letter(len(headers)) + str(max(ws.max_row, 1))
    return ws


wb = Workbook()
wb.remove(wb.active)

# --- sheet 1: one row per AP ---
ap_headers = ["Site", "AP name", "AP MAC", "Model", "Serial", "Firmware", "Status", "Uptime (days)", "Device Profile"]
ap_widths = [28, 22, 15, 8, 16, 12, 12, 9, 40]
for band in BANDS:
    b = BAND_NAMES[band]
    ap_headers = ap_headers + [b + " channel", b + " width (MHz)", b + " power (dBm)", b + " ant gain (dBi)",
                               b + " clients", b + " SSIDs", b + " SSIDs (config)", b + " SSIDs (live)"]
    ap_widths = ap_widths + [9, 8, 8, 8, 8, 60, 9, 9]
ap_headers = ap_headers + ["SSID check (config vs live)", "Outside Device Profile limits", "Power below Device Profile min (info)"]
ap_widths = ap_widths + [40, 50, 50]

ap_rows.sort(key=lambda r: (str(r[0]).lower(), str(r[1]).lower()))
ws = add_sheet(wb, "AP SSIDs by Channel", ap_headers, ap_rows, ap_widths)
# Colour the band headers so the three blocks are easy to find
for i, band in enumerate(BANDS):
    for c in range(10 + i * 8, 18 + i * 8):
        ws.cell(row=1, column=c).fill = band_fills[band]
check_col = len(ap_headers) - 2
for row_num in range(2, ws.max_row + 1):
    check = ws.cell(row=row_num, column=check_col).value
    if check == "AP not connected":
        ws.cell(row=row_num, column=7).fill = grey_fill
    elif check not in ("OK", None) and not str(check).startswith("Not checked"):
        ws.cell(row=row_num, column=check_col).fill = amber_fill
    if ws.cell(row=row_num, column=check_col + 1).value:
        ws.cell(row=row_num, column=check_col + 1).fill = red_fill

# --- sheet 2: one row per AP / band / SSID ---
detail_headers = ["Site", "AP name", "AP MAC", "Model", "Status", "Device Profile", "Band", "Channel",
                  "Width (MHz)", "Power (dBm)", "Radio on", "SSID", "WLAN ID", "Comes from", "Auth",
                  "VLAN", "Hidden", "Note"]
detail_rows.sort(key=lambda r: (str(r[0]).lower(), str(r[1]).lower(), BANDS.index([k for k, v in BAND_NAMES.items() if v == r[6]][0]), str(r[11]).lower()))
add_sheet(wb, "AP SSID Detail", detail_headers, detail_rows,
          [28, 22, 15, 8, 12, 40, 8, 9, 8, 8, 10, 24, 38, 40, 10, 10, 8, 24])

# --- sheet 3: APs that need a look ---
check_rows.sort(key=lambda r: (str(r[0]).lower(), str(r[1]).lower()))
add_sheet(wb, "Needs a Look", ["Site", "AP name", "AP MAC", "Model", "Device Profile",
                               "SSID check (config vs live)", "Outside Device Profile limits"],
          check_rows, [28, 22, 15, 8, 40, 50, 60])

# --- sheet 4: Device Profiles ---
templates_for_dp = {}
for t in templates:
    if t.get("filter_by_deviceprofile"):
        for dp_id in t.get("deviceprofile_ids") or []:
            templates_for_dp.setdefault(dp_id, []).append(t.get("name", ""))
dp_rows = []
for dp_id in sorted(dp_ap_count.keys(), key=lambda d: deviceprofiles_by_id.get(d, {}).get("name", "(none)").lower()):
    dp = deviceprofiles_by_id.get(dp_id, {})
    rc = dp.get("radio_config") or {}
    r = [dp.get("name", "(none)" if dp_id == "" else dp_id), dp_id, dp_ap_count[dp_id]]
    for band in BANDS:
        bc = rc.get("band_" + band) or {}
        if bc.get("disabled"):
            r = r + ["disabled", "", "", ""]
        else:
            pr = ""
            if bc.get("power_min") is not None or bc.get("power_max") is not None:
                pr = str(bc.get("power_min", "")) + "-" + str(bc.get("power_max", ""))
            r = r + ["on" if bc else "", ", ".join([str(c) for c in bc.get("channels") or []]), pr, bc.get("bandwidth", "")]
    r = r + [", ".join(sorted(templates_for_dp.get(dp_id, []), key=str.lower))]
    dp_rows.append(r)
dp_headers = ["Device Profile", "Device Profile ID", "APs"]
for band in BANDS:
    b = BAND_NAMES[band]
    dp_headers = dp_headers + [b + " radio", b + " channels", b + " power range", b + " width"]
dp_headers = dp_headers + ["WLAN Templates filtered to this profile"]
add_sheet(wb, "Device Profiles", dp_headers, dp_rows,
          [44, 38, 7, 8, 14, 10, 8, 8, 40, 10, 8, 8, 30, 10, 8, 60])

# --- sheet 5: summary ---
ws2 = wb.create_sheet("Summary", 0)
ws2.append(["AP SSID / Channel / Device Profile report"])
ws2["A1"].font = title_font
ws2.append([])
summary_rows = [
    ["Run time", run_time],
    ["Org", org_name],
    ["Org ID", ORG_ID],
    ["Site filter", SITE_NAME_FILTER if SITE_NAME_FILTER != "" else "(all sites)"],
    ["Sites checked", len(sites)],
    ["APs", total_aps],
    ["APs connected", aps_connected],
    ["Device Profiles in use", len(dp_ap_count)],
    ["AP / band / SSID rows", len(detail_rows)],
    [],
    ["Connected APs where config and live SSID count differ", aps_ssid_mismatch],
    ["Connected APs outside their Device Profile limits (channel not allowed / power above max)", aps_outside_limits],
    ["Connected APs with power below Device Profile min (info - EIRP cap)", aps_power_capped],
    ["APs with WxLAN-tag WLANs (not checked)", aps_wxtag],
    [],
    ["API calls", api_calls],
    [],
    ["How SSIDs are worked out", "WLANs from /sites/{id}/wlans/derived, kept if: enabled; the WLAN Template's "
        "Device Profile filter (if set) includes the AP's Device Profile; and 'apply to specific APs' (if set) "
        "includes the AP. Each WLAN goes on the bands in its 'bands' setting."],
    ["SSIDs (live)", "num_wlans from the AP's live radio stats - how many SSIDs the radio is really broadcasting. "
        "Amber in 'SSID check' = the worked-out list doesn't match."],
    ["Outside Device Profile limits", "Live channel not in the allowed channel list, or live power above "
        "power_max, from the AP's effective radio config (Device Profile plus any AP-level override). "
        "Blank if no limits are set."],
    ["Power below Device Profile min (info)", "Not flagged as a fault: Mist lowers the power below power_min to "
        "stay within the legal EIRP limit, e.g. with a high-gain (directional) antenna or on a low-power channel."],
    ["Colour key", "Grey status = AP not connected; Amber = SSID count mismatch; Red = outside Device Profile limits"],
]
for r in summary_rows:
    ws2.append(r)
for row in ws2.iter_rows(min_row=3, max_col=1):
    for cell in row:
        cell.font = Font(bold=True)
for row in ws2.iter_rows(min_row=3, min_col=2, max_col=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
ws2.column_dimensions["A"].width = 60
ws2.column_dimensions["B"].width = 100

os.makedirs(REPORT_FOLDER, exist_ok=True)
wb.save(report_file)

# ---------- 4. summary ----------

print("")
print("==== Summary ====")
print("Sites checked:                    " + str(len(sites)))
print("APs:                              " + str(total_aps) + " (" + str(aps_connected) + " connected)")
print("Device Profiles in use:           " + str(len(dp_ap_count)))
print("AP / band / SSID rows:            " + str(len(detail_rows)))
print("SSID count mismatch (connected):  " + str(aps_ssid_mismatch))
print("Outside Device Profile limits:    " + str(aps_outside_limits))
print("Power below DP min (info only):   " + str(aps_power_capped))
print("APs with WxLAN-tag WLANs:         " + str(aps_wxtag))
print("API calls:                        " + str(api_calls))
print("Excel report:                     " + report_file)
